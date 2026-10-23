"""
export.py - the Word document, the app's final output.

The document is meant to be a fully usable work product for the lawyer: easy to edit in Word, with everything that
points to the AI in an Annex that can be ignored.
    main part   the lawyer's work product: only the lawyer's final choices (no "by you" tags, no struck-through tool
                labels, nothing dismissed), Word's own styles, borders and bullets, Times New Roman throughout; colour
                only on the two relationship labels (light shading, changeable in Word); the Commission's result in
                plain bold
    Annex       how the overview was prepared (the tool, its limits, the run summary), the pairs without an overlap,
                the passages behind an existing relationship, the activities not compared
The content comes from app.report_model(); this module only lays it out (python-docx). No AI call.
"""
import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

FONT = "Times New Roman"
INK = "22313F"
GREY = "595959"
LABEL = {"h": ("Potential horizontal overlap", "FDECEA", "8A1C12"),
         "v": ("Potential vertical relationship", "FFF3D1", "6E4B00"),
         "q": ("To be checked", "ECEAE5", "4A463F")}
PREFIX_SEP = ":  "        # app.PREFIX_SEP: the products a bullet concerns (set in italics)
KEY = [("Confirmed", "both products fall in a market the Commission defined"),
       ("Possible", "both may fall in the same market, depending on the product format"),
       ("Not excluded", "the Commission left the market definition open"),
       ("Not supported", "the Commission puts the products in separate markets"),
       ("No precedent", "no decision in the database covers the products (or, for a shop or wholesaler, its retail or "
                        "wholesale market)"),
       ("Markets defined / One side defined", "vertical rows: the decisions define the market of both sides' products "
                                              "/ of one side's only; not the link between the two companies")]
HOW = [
    ("Research", "Each company was researched on its own, in two searches of public websites (its main business, and "
                 "what such a search easily misses: smaller brands, sales to other businesses, own shops and cafés, "
                 "side lines). Each search saw only that company, never another company of this overview or the "
                 "planned transaction. The app checked every quotation the search relied on against the page."),
    ("Preliminary assessment", "Every activity of the target was compared with every activity of each portfolio "
                               "company, one pair at a time, by an AI model that saw only two anonymised "
                               "descriptions. The model judged whether the products are the same or similar for the "
                               "customers, and whether one could be an input for the other or be resold by it. The app "
                               "itself checked the level of trade, the direct customers and the conditions for supply "
                               "and resale. The assessment is deliberately generous; the reviewing lawyer decided "
                               "which relationships to keep."),
    ("Commission market definitions", "The products of each row were searched in a database of Commission decisions "
                                      "and the Market Definition Notice. The Commission's sentences were quoted word "
                                      "for word and checked against the decisions. A decision outside the database is "
                                      "not found; “No precedent” does not mean that there is no market. For a shop or "
                                      "a wholesaler, the decisions describe the markets of the products it sells, not "
                                      "its own retail or wholesale market."),
    ("Existing relationships", "Found by the app's code, by matching the names of the other companies in each "
                               "company's own research, never by asking the AI about two companies together; the "
                               "nature of the relationship is not assessed."),
    ("Not covered", "Geography and market shares; anything the companies do not publish."),
]


# =============================================================== helpers
def _font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    f = rpr.find(qn("w:rFonts"))
    if f is None:
        f = OxmlElement("w:rFonts")
        rpr.insert(0, f)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        f.set(qn(a), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _shd(pr, fill):
    s = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        s.set(qn(k), v)
    pr.append(s)


def _borders(table, color="BFBFBF"):
    """Word's own table borders (table level), changeable with the Borders menu."""
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", color)):
            el.set(qn(k), v)
        b.append(el)
    table._element.tblPr.append(b)


def _margins(table, tb=50, lr=90):
    m = OxmlElement("w:tblCellMar")
    for edge, v in (("top", tb), ("left", lr), ("bottom", tb), ("right", lr)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    table._element.tblPr.append(m)


def _row_props(row, header=False):
    pr = row._tr.get_or_add_trPr()
    if header:
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        pr.append(el)
    else:
        pr.append(OxmlElement("w:cantSplit"))


def _widths(table, cms):
    table.autofit = False
    for row in table.rows:
        for c, w in zip(row.cells, cms):
            c.width = Cm(w)
    for gc, w in zip(table._element.find(qn("w:tblGrid")).findall(qn("w:gridCol")), cms):
        gc.set(qn("w:w"), str(int(w * 567)))


def _fmt(p, after=2, before=0, keep=False):
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), 1.05
    if keep:
        pf.keep_with_next = True
    return p


def _para(doc, text="", size=11, bold=False, italic=False, color=None, after=3, before=0, keep=False, style=None):
    p = _fmt(doc.add_paragraph(style=style) if style else doc.add_paragraph(), after, before, keep)
    if text:
        _font(p.add_run(text), size, bold, italic, color)
    return p


class _Cell:
    """Paragraphs in a table cell, the first one into the cell's own paragraph."""

    def __init__(self, cell):
        self.c, self.used = cell, False

    def p(self, style=None):
        if not self.used:
            self.used = True
            p = self.c.paragraphs[0]
            if style:
                p.style = self.c.part.document.styles[style]
        else:
            p = self.c.add_paragraph(style=style) if style else self.c.add_paragraph()
        return _fmt(p, after=2)


def _label(p, kind, size=10):
    text, fill, col = LABEL[kind]
    r = _font(p.add_run(text), size, True, False, col)
    _shd(r._element.get_or_add_rPr(), fill)


def _heading(doc, text, level):
    p = _fmt(doc.add_paragraph(style=f"Heading {level}"), after=4, before=10 if level == 1 else 8, keep=True)
    _font(p.add_run(text), 14 if level == 1 else 12, True, False, "000000")
    return p


def _head(table, names):
    for c, t in zip(table.rows[0].cells, names):
        p = _Cell(c).p()
        for k, part in enumerate(t.split("*")):        # *word*: in italics (the defined terms target, portfolio)
            if part:
                _font(p.add_run(part), 10, True, k % 2 == 1, "FFFFFF")
        _shd(c._element.get_or_add_tcPr(), INK)
    _row_props(table.rows[0], header=True)


def _field(p, instr):
    r = _font(p.add_run(), 9, color=GREY)
    for kind, txt in (("begin", None), (None, instr), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        r._element.append(e)


def _bullet(cell_or_doc, text, size=10):
    """A Word bullet; the products a bullet concerns (before PREFIX_SEP) in italics."""
    p = cell_or_doc.p(style="List Bullet") if isinstance(cell_or_doc, _Cell) else \
        _fmt(cell_or_doc.add_paragraph(style="List Bullet"), after=1)
    head, sep, rest = text.partition(PREFIX_SEP)
    if sep:
        _font(p.add_run(head + ": "), size, False, True)
        _font(p.add_run(rest[:1].upper() + rest[1:] if not rest[:1].isupper() else rest), size)
    else:
        _font(p.add_run(text[:1].upper() + text[1:]), size)
    return p


def _styles(doc):
    for name, size in (("Normal", 11), ("Heading 1", 14), ("Heading 2", 12), ("List Bullet", 10)):
        st = doc.styles[name]
        st.font.name, st.font.size = FONT, Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
    doc.styles["List Bullet"].paragraph_format.space_after = Pt(1)


def _page(doc, target):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.6)
    hp = sec.header.paragraphs[0]
    for pos in (4680, 9360):
        hp.paragraph_format.tab_stops.add_tab_stop(pos * 635, WD_TAB_ALIGNMENT.CLEAR)
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(25.7), WD_TAB_ALIGNMENT.RIGHT)
    _font(hp.add_run(f"Preliminary overlap overview · {target}"), 9, color=GREY)
    hp.add_run("\t")
    _font(hp.add_run("Draft · privileged and confidential"), 9, italic=True, color=GREY)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(fp.add_run("Page "), 9, color=GREY)
    _field(fp, "PAGE")
    _font(fp.add_run(" of "), 9, color=GREY)
    _field(fp, "NUMPAGES")


def _activity(cell, act, company=None):
    w = _Cell(cell)
    _font(w.p().add_run(act["name"]), 10, True)
    src = "".join(f"[{n}]" for n in act["sources"])
    _font(w.p().add_run(f"{act['stage']} {src}".strip()), 9, color=GREY)
    items = ", ".join(act["items"][:6]) + (" …" if len(act["items"]) > 6 else "")
    if items:
        _font(w.p().add_run(items), 9, color=GREY)


def _company_table(doc, model, company, blocks):
    step4 = model["step4"]
    cols = [f"{model['target']} Activity", f"{company} Activity", "Preliminary Assessment"]
    cols += (["Commission Market Definition"] if step4 else []) + ["Notes"]
    t = doc.add_table(rows=1, cols=len(cols))
    _borders(t)
    _margins(t)
    _head(t, cols)
    for block in blocks:
        start = len(t.rows)
        for i, r in enumerate(block["rows"]):
            row = t.add_row()
            _row_props(row)
            c = row.cells
            if i == 0:
                _activity(c[0], block["target"])
            _activity(c[1], r["b_rec"])
            w = _Cell(c[2])
            for lab in r["labels"]:
                _label(w.p(), lab["kind"])
                if lab["basis"]:
                    _font(w.p().add_run(lab["basis"]), 9.5, color=GREY)
                if lab["regard"]:
                    _font(w.p().add_run("With regard to"), 10)
                    for who, what in lab["regard"]:
                        p = w.p()
                        p.paragraph_format.left_indent = Cm(0.3)
                        _font(p.add_run(who + ": "), 10, False, True)
                        _font(p.add_run(what), 10)
                for line in lab["lines"]:
                    _font(w.p().add_run(line[:1].upper() + line[1:]), 10)
            for line in r["edited"]:
                _font(w.p().add_run(line), 10)
            if r.get("relationship"):
                p = w.p()
                _font(p.add_run("Existing relationship: "), 9.5, True)
                _font(p.add_run(f"the public sources of {model['target']} and {company} mention each other "
                                "(Annex C)."), 9.5)
            if step4:
                w = _Cell(c[3])
                for label, bullets, outcome in r["commission"]:
                    name = _label_name(label) if label and len(r["commission"]) > 1 else ""
                    # "Possible (horizontal label)", or the label alone where there is no result (e.g. a failed search)
                    head = f"{outcome} ({name})" if outcome and name else outcome or name[:1].upper() + name[1:]
                    if head:
                        _font(w.p().add_run(head + ":"), 10, True)
                    for b in bullets:
                        _bullet(w, b)
                if not w.used:
                    w.p()
            w = _Cell(c[4 if step4 else 3])
            views = list(dict.fromkeys(r["views"]))
            if views:
                _font(w.p().add_run(" · ".join(v[:1].upper() + v[1:] for v in views)), 10, True)
            if r["note"]:
                _font(w.p().add_run(r["note"]), 10)
            if not w.used:
                w.p()
        end = len(t.rows) - 1
        if end > start:
            t.cell(start, 0).merge(t.cell(end, 0))
        t.cell(start, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _widths(t, [4.2, 4.2, 6.9, 6.9, 3.5] if step4 else [5.0, 5.0, 11.2, 4.5])


# =============================================================== the document
def _label_name(label):
    """The label of step ③ a part of the Commission cell belongs to (as on screen)."""
    if label == "vertical, both directions":
        return "vertical labels, both directions"
    if label.startswith("vertical, "):
        return "vertical label, " + label[len("vertical, "):]
    return "horizontal label" if label.startswith("horizontal") else label


def _markets_table(doc, model, n):
    """Every market a product of the kept rows falls in, once, with the paragraph it rests on and each company's
    products in it (as the table at the top of step ④)."""
    _heading(doc, f"{n}. Markets in the Commission's Decisions", 1)
    _para(doc, "Every market a product of the rows in section 2 falls in, once, with the paragraph it rests on. The "
               "decisions define markets, not the relationship between the companies.", 10.5, after=4)
    t = doc.add_table(rows=1, cols=4)
    _borders(t)
    _margins(t)
    _head(t, ["Market", "Commission's Decision", model["target"], "Portfolio Companies"])
    group = None
    for x in model["markets"]:
        if x["group"] != group:
            group = x["group"]
            row = t.add_row()
            _row_props(row)
            cell = row.cells[0].merge(row.cells[3])
            _shd(cell._tc.get_or_add_tcPr(), "F2EFE8")
            _font(cell.paragraphs[0].add_run(group), 10, True)
        row = t.add_row()
        _row_props(row)
        _font(_Cell(row.cells[0]).p().add_run(x["market"]), 10, True)
        for c, lines in ((row.cells[1], x["decisions"]), (row.cells[2], x["target"]), (row.cells[3], x["portcos"])):
            w = _Cell(c)
            for line in lines or ["–"]:
                _font(w.p().add_run(line), 9.5)
    _widths(t, [4.2, 5.4, 8.0, 8.1])


def docx_bytes(model):
    doc = Document()
    _styles(doc)
    _page(doc, model["target"])
    target = model["target"]

    _para(doc, "Preliminary Overlap Overview", 20, True, after=6)
    for k, v in (("Target: ", f"{target} (the business to be acquired)"),
                 ("Portfolio companies: ", ", ".join(model["portcos"])), ("Date: ", model["date"])):
        p = _para(doc, after=1)
        _font(p.add_run(k), 11, True)
        _font(p.add_run(v), 11)
    _para(doc, "Draft for internal review. Product markets only: geography and market shares, which decide whether a "
               "market is “affected”, are not assessed."
          + (f" The Commission decisions cited describe each market as it was at the time ({model['years']})."
             if model["years"] else ""), 10.5, before=5, after=2)
    _para(doc, "Prepared with the help of an AI-assisted screening tool; see the Annex.", 9.5, italic=True,
          color=GREY, after=8)

    # 1 summary
    _heading(doc, "1. Summary", 1)
    t = doc.add_table(rows=1, cols=3)
    _borders(t)
    _margins(t)
    _head(t, ["Portfolio Company", "Potential Horizontal Overlaps", "Potential Vertical Relationships"])
    for s in model["summary"]:
        row = t.add_row()
        _row_props(row)
        _font(_Cell(row.cells[0]).p().add_run(s["portco"]), 10.5, True)
        for c, entries in ((row.cells[1], s["horizontal"]), (row.cells[2], s["vertical"])):
            w = _Cell(c)
            if not entries:
                _font(w.p().add_run("None"), 10)
            for pair, res, view in entries:
                p = w.p(style="List Bullet")
                _font(p.add_run(pair + " "), 10)
                bits = [x for x in (res, view) if x]
                if bits:
                    _font(p.add_run(f"({'; '.join(bits)})"), 10, False, True)
    _widths(t, [4.4, 11.2, 10.1])
    _para(doc, "Target activity first; the arrow points to the customer.", 9.5, italic=True, color=GREY, before=3,
          after=6)
    if model["step4"]:
        _para(doc, "Commission Results Used in This Overview", 10, True, before=4, after=1, keep=True)
        for res, text in KEY:
            p = _para(doc, after=0)
            p.paragraph_format.left_indent = Cm(0.4)
            _font(p.add_run(res + ": "), 9.5, True)
            _font(p.add_run(text), 9.5)

    # 2 the portfolio companies
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _heading(doc, "2. The Portfolio Companies", 1)
    for i, sec in enumerate(model["sections"], start=1):
        _heading(doc, f"2.{i} {sec['portco']}", 2)
        _para(doc, sec["sentence"], 11, after=4, keep=True)
        if sec.get("relationship"):
            p = _para(doc, after=4)
            _font(p.add_run("Existing relationship: "), 10.5, True)
            _font(p.add_run(f"the public sources of {target} and {sec['portco']} mention each other (Annex C)."), 10.5)
        if sec["blocks"]:
            _company_table(doc, model, sec["portco"], sec["blocks"])
        elif not sec["own"]:
            _para(doc, "The pairs compared are listed in Annex B.", 10.5, color=GREY, after=6)
        if sec["own"]:
            _para(doc, "Own Use (production for the group itself, which can only be the customer side)", 10.5, True,
                  before=8, after=2, keep=True)
            _company_table(doc, model, sec["portco"], sec["own"])

    # 3 the markets in the Commission's decisions (step ④), each once; the rows above name them
    n = 3
    if model.get("markets"):
        _markets_table(doc, model, n)
        n += 1

    # open points
    if model["open_points"]:
        _heading(doc, f"{n}. Open Points", 1)
        groups = {}
        for o in model["open_points"]:
            groups.setdefault(o["portco"], []).append(o)
        for portco, items in groups.items():
            _para(doc, f"{target} × {portco}", 11, True, after=1, keep=True)
            for o in items:
                p = _fmt(doc.add_paragraph(style="List Bullet"), after=1)
                _font(p.add_run(o["text"] + " "), 11)
                _font(p.add_run(f"({o['pair']})"), 10, False, True, GREY)
        n += 1

    # sources and decisions
    _heading(doc, f"{n}. Sources", 1)
    for src in model["sources"]:
        _para(doc, src["company"], 11, True, before=3, after=1, keep=True)
        for num, title, url in src["items"]:
            p = _para(doc, after=0)
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.8), Cm(-0.8)
            _font(p.add_run(f"[{num}]\t"), 10.5)
            _font(p.add_run(title), 10.5, False, True)
            if url:                         # a description typed in by the user has no address and no access date
                _font(p.add_run(f", {url}" + (f", accessed {src['date']}" if src["date"] else "")), 10.5)
    if model["decisions"]:
        _para(doc, "Commission Decisions and Notice Cited", 11, True, before=6, after=1, keep=True)
        for d in model["decisions"]:
            p = _para(doc, after=0)
            p.paragraph_format.left_indent = Cm(0.8)
            name = d["title"] if d["case"] == "Notice" else f"{d['case']} {d['title']}"
            _font(p.add_run(name), 10.5, False, True)
            paras = ", ".join(f"({x})" for x in d["paras"])
            _font(p.add_run((f", {d['date']}" if d["date"] else "") + (f"; paragraphs cited: {paras}" if paras
                                                                        else "")), 10.5)

    # Annex
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _heading(doc, "Annex: How This Overview Was Prepared", 1)
    _para(doc, "This annex records how the screening tool prepared the overview. It is not part of the assessment and "
               "can be removed.", 10.5, italic=True, color=GREY, after=6)
    _heading(doc, "A. How the Tool Works and Its Limits", 2)
    for title, text in HOW:
        p = _para(doc, after=3)
        _font(p.add_run(title + ". "), 10.5, True)
        _font(p.add_run(text), 10.5)
    p = _para(doc, before=4, after=6)
    _font(p.add_run("Run summary. "), 10.5, True)
    _font(p.add_run(model["annex"]["run_summary"]), 10.5)

    _heading(doc, "B. Pairs Compared Without an Overlap", 2)
    pairs = model["annex"]["pairs"]
    if pairs:
        _para(doc, "Every pair not shown in section 2, with the reason. In the tool's reasons, “A” is the target's "
                   "activity and “B” the portfolio company's activity (the AI saw them without names).", 10.5,
              color=GREY, after=3)
        t = doc.add_table(rows=1, cols=3)
        _borders(t)
        _margins(t)
        _head(t, ["Portfolio Company", "Pair (*target* activity × *portfolio* activity)", "Reason"])
        for x in pairs:
            row = t.add_row()
            _row_props(row)
            for c, txt in zip(row.cells, (x["portco"], x["pair"], x["reason"])):
                _font(_Cell(c).p().add_run(txt), 9.5)
        _widths(t, [3.6, 10.6, 11.5])
    else:
        _para(doc, "None.", 10.5)

    _heading(doc, "C. Existing Relationships", 2)
    rels = model["annex"]["relationships"]
    if not rels:
        _para(doc, "None found.", 10.5)
    for rel in rels:
        _para(doc, f"{target} and {rel['portco']}", 10.5, True, after=1, keep=True)
        for who, passage in rel["passages"]:
            t_ = passage.strip()
            quoted = t_ if t_[:1] in "“\"„«" else f"“{t_}”"
            _para(doc, f"{who}'s sources: {quoted}", 10.5, after=2)

    _heading(doc, "D. Activities Not Compared", 2)
    for line in model["annex"]["not_compared"] or ["None."]:
        _para(doc, line, 10.5, after=1)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()